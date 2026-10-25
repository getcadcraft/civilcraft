# SPDX-License-Identifier: Apache-2.0
"""cobie_xlsx.py — openpyxl as the independent oracle for BimCraft's COBie workbook (0.7 LC.12, session bim-ops).

    python -I cobie_xlsx.py read <file.xlsx>     → JSON: {sheets: {name: [[value…]…]}, quoted: [sheet!cell…], formulas: n}
    python -I cobie_xlsx.py write <file.xlsx>    a workbook written the way a spreadsheet program does: shared strings,
                                                 a formula, a real date cell, numbers, a boolean, COBie-shaped sheets
Run with tools/takeoffcraft-test/venv/bin/python (openpyxl 3.1). Shares no code with assets-bim-xlsx.js.
"""
import datetime
import json
import sys

import openpyxl


def read(path):
    wb = openpyxl.load_workbook(path)
    out, quoted, formulas = {}, [], 0
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            vals = []
            for c in row:
                if c.data_type == 'f':
                    formulas += 1
                if c.quotePrefix:
                    quoted.append(ws.title + '!' + c.coordinate)
                v = c.value
                vals.append(v if isinstance(v, (int, float, str)) or v is None else str(v))
            rows.append(vals)
        out[ws.title] = rows
    print(json.dumps({'sheets': out, 'quoted': quoted, 'formulas': formulas}))


def write(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Component'
    ws.append(['Name', 'CreatedBy', 'CreatedOn', 'TypeName', 'Space', 'Description', 'ExtSystem', 'ExtObject', 'ExtIdentifier',
               'SerialNumber', 'InstallationDate', 'WarrantyStartDate', 'TagNumber', 'BarCode', 'AssetIdentifier'])
    # matched by name (no ExtIdentifier); dates as real date cells; the serial typed as a number
    ws.append(['F-101', 'fm@example.com', datetime.datetime(2026, 10, 1, 9, 0), 'Desk type', 'n/a', 'Desk', 'Other CAFM', 'Component', 'n/a',
               4471, datetime.date(2025, 3, 1), datetime.date(2025, 3, 1), 'F-101', 'BC-9', 'F-101'])
    ws.append(['Not in this model', 'fm@example.com', 'n/a', 'n/a', 'n/a', 'n/a', 'n/a', 'n/a', 'n/a', 'X', 'n/a', 'n/a', 'n/a', 'n/a', 'n/a'])
    ws['P1'] = 'Check'
    ws['P2'] = '=LEN(A2)'          # a formula: openpyxl writes it without a cached value
    ws['Q1'] = 'Ok'
    ws['Q2'] = True
    t = wb.create_sheet('Type')
    t.append(['Name', 'CreatedBy', 'CreatedOn', 'Category', 'Description', 'AssetType', 'Manufacturer', 'ModelNumber', 'WarrantyGuarantorParts',
              'WarrantyDurationParts', 'WarrantyGuarantorLabor', 'WarrantyDurationLabor', 'WarrantyDurationUnit', 'ExtSystem', 'ExtObject', 'ExtIdentifier',
              'ReplacementCost', 'ExpectedLife', 'DurationUnit'])
    # the warranty five years from its start; life 180 months = 15 years
    t.append(['Desk type', 'fm@example.com', 'n/a', 'Furniture', 'n/a', 'Fixed', 'Steelcase', 'Ology', 'n/a', 5, 'n/a', 'n/a', 'Year', 'n/a', 'n/a', 'n/a', 'n/a', 180, 'Month'])
    wb.save(path)


if __name__ == '__main__':
    {'read': read, 'write': write}[sys.argv[1]](sys.argv[2])
