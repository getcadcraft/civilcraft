# SPDX-License-Identifier: Apache-2.0
"""xlsx_check.py — an independent reader for the workbook TakeoffCraft writes.

    python3 xlsx_check.py <file.xlsx>

Prints JSON: the sheet names, every cell of every sheet (value or formula), the cached values,
and the Estimate sheet RE-CALCULATED here from its own formulas and inputs with Python's
Decimal (ROUND half away from zero, like a spreadsheet), so the Node test can compare the app's
numbers with an implementation that shares no code with it. Uses openpyxl when it is installed
(tools/takeoffcraft-test/venv), else only the standard library's zipfile + ElementTree.
"""
import json
import re
import sys
import zipfile
from decimal import Decimal, ROUND_HALF_UP
import xml.etree.ElementTree as ET

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def read_stdlib(path):
    z = zipfile.ZipFile(path)
    bad = z.testzip()
    if bad:
        raise SystemExit('corrupt member ' + bad)
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    names = [s.get('name') for s in wb.find('m:sheets', NS)]
    sheets = {}
    for i, name in enumerate(names):
        root = ET.fromstring(z.read('xl/worksheets/sheet%d.xml' % (i + 1)))
        cells, cached = {}, {}
        for c in root.iter('{%s}c' % NS['m']):
            r = c.get('r')
            f = c.find('m:f', NS)
            v = c.find('m:v', NS)
            t = c.get('t')
            if t == 'inlineStr':
                cells[r] = ''.join(x.text or '' for x in c.iter('{%s}t' % NS['m']))
                cached[r] = cells[r]
            elif f is not None:
                cells[r] = '=' + (f.text or '')
                cached[r] = float(v.text) if v is not None else None
            elif v is not None:
                cells[r] = float(v.text)
                cached[r] = cells[r]
        sheets[name] = {'cells': cells, 'cached': cached}
    return names, sheets


def read_openpyxl(path):
    import openpyxl
    wf = openpyxl.load_workbook(path)                  # formulas
    wv = openpyxl.load_workbook(path, data_only=True)  # cached values
    sheets = {}
    for name in wf.sheetnames:
        cells, cached = {}, {}
        for row in wf[name].iter_rows():
            for c in row:
                if c.value is None:
                    continue
                cells[c.coordinate] = c.value
                cached[c.coordinate] = wv[name][c.coordinate].value
        sheets[name] = {'cells': cells, 'cached': cached}
    return wf.sheetnames, sheets


def D(v):
    return Decimal(repr(float(v)))


def recalc(cells):
    """Evaluate a sheet's formulas (the Estimate, the Bid leveling): X*(1+Y), ROUND(X*Y,2), X, X+Y(+Z…), X-Y, SUM(Xa:Xb)."""
    vals = {}

    def get(ref):
        if ref in vals:
            return vals[ref]
        v = cells.get(ref)
        if v is None or v == '':
            return Decimal(0)
        if isinstance(v, str) and v.startswith('='):
            vals[ref] = ev(v[1:])
            return vals[ref]
        if isinstance(v, (int, float)):
            return D(v)
        return Decimal(0)

    def ev(f):
        m = re.fullmatch(r'([A-Z]+\d+)\*\(1\+([A-Z]+\d+)\)', f)
        if m:
            return get(m[1]) * (1 + get(m[2]))
        m = re.fullmatch(r'ROUND\(([A-Z]+\d+)\*([A-Z]+\d+),2\)', f)
        if m:
            return (get(m[1]) * get(m[2])).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        m = re.fullmatch(r'([A-Z]+\d+)-([A-Z]+\d+)', f)   # a difference (the Bid leveling sheet, LC.13)
        if m:
            return get(m[1]) - get(m[2])
        if re.fullmatch(r'[A-Z]+\d+(\+[A-Z]+\d+)*', f):      # a reference, or a sum of references (the markups block, TK.4)
            return sum((get(r) for r in f.split('+')), Decimal(0))
        m = re.fullmatch(r'SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)', f)
        if m:
            return sum((get(m[1] + str(r)) for r in range(int(m[2]), int(m[4]) + 1)), Decimal(0))
        raise SystemExit('unexpected formula ' + f)

    out = {}
    for ref, v in cells.items():
        if isinstance(v, str) and v.startswith('='):
            out[ref] = float(get(ref))
    return out


def main():
    path = sys.argv[1]
    try:
        names, sheets = read_openpyxl(path)
        reader = 'openpyxl'
    except ImportError:
        names, sheets = read_stdlib(path)
        reader = 'zipfile'
    zipfile.ZipFile(path).testzip()
    est = sheets.get('Estimate', {}).get('cells', {})
    other = {n: recalc(sheets[n]['cells']) for n in names if n == 'Bid leveling'}   # LC.13: its formulas too
    print(json.dumps({'reader': reader, 'names': names, 'sheets': sheets, 'recalc': recalc(est), 'recalcOther': other}, default=str))


main()
